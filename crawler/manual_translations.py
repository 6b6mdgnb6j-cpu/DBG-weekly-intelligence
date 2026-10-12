import hashlib
import json
import os
import re
import tempfile
from copy import copy
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from crawler.translation_quality import FAIL, validate_translation_field


SCHEMA_VERSION = 1
VISIBLE_SHEET = "Manual Review"
METADATA_SHEET = "_metadata"
VISIBLE_HEADERS = (
    "Item ID",
    "Source",
    "Field",
    "English Original",
    "Current CN",
    "QA Failure",
    "Manual Translation",
)
METADATA_HEADERS = (
    "schema_version",
    "week_start",
    "visible_row",
    "item_id",
    "source",
    "date",
    "url",
    "field",
    "original_text",
    "original_sha256",
    "title_original",
)

INVESTMENT_ASSOCIATION_SOURCES = {"英国投资管理协会", "Investment Association"}
INVESTMENT_ASSOCIATION_HOST = "www.theia.org"


class ManualTranslationError(ValueError):
    pass


def original_sha256(text):
    return hashlib.sha256((text or "").encode("utf-8")).hexdigest()


def stable_item_id(item, field, original_text):
    payload = {
        "source": item.get("source") or "",
        "date": str(item.get("date") or ""),
        "url": item.get("url") or "",
        "field": field,
        "original_sha256": original_sha256(original_text),
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _investment_association_canonical_url(url):
    parsed = urlsplit(url or "")
    if parsed.scheme.lower() != "https" or parsed.netloc.lower() != INVESTMENT_ASSOCIATION_HOST:
        return ""
    path = re.sub(r"%2e", ".", parsed.path, flags=re.I)
    if path == "/index.php":
        path = "/"
    elif path.startswith("/index.php/"):
        path = path[len("/index.php"):]
    path = path.rstrip("/") or "/"
    return urlunsplit(("https", INVESTMENT_ASSOCIATION_HOST, path, parsed.query, ""))


def _approved_override_for_item(overrides, item, field, original_text):
    item_id = stable_item_id(item, field, original_text)
    exact = overrides.get(item_id)
    if exact:
        return item_id, exact

    source = item.get("source") or ""
    if source not in INVESTMENT_ASSOCIATION_SOURCES:
        return item_id, None
    canonical_url = _investment_association_canonical_url(item.get("url") or "")
    if not canonical_url:
        return item_id, None

    expected_hash = original_sha256(original_text)
    candidates = [
        (override_id, override)
        for override_id, override in overrides.items()
        if override.get("source") == source
        and str(override.get("date") or "") == str(item.get("date") or "")
        and override.get("field") == field
        and override.get("original_sha256") == expected_hash
        and _investment_association_canonical_url(override.get("url") or "") == canonical_url
    ]
    if len(candidates) != 1:
        return item_id, None
    return candidates[0]


def _field_values(item, field):
    if field == "title":
        return item.get("title_original") or item.get("title") or "", item.get("title_cn") or ""
    return item.get("summary_original") or item.get("summary") or "", item.get("summary_cn") or ""


def build_queue_entries(weekly_payload, qa_report):
    items = weekly_payload.get("items") or []
    grouped = {}
    for issue in qa_report.get("content_issues") or []:
        issue_id = issue.get("issue_id") or issue.get("type") or ""
        field = issue.get("field") or ""
        if issue.get("severity") != FAIL or field not in {"title", "summary"}:
            continue
        if not issue_id.startswith("translation_"):
            continue
        index = issue.get("index")
        if not isinstance(index, int) or index < 0 or index >= len(items):
            continue
        grouped.setdefault((index, field), []).append(issue)

    entries = []
    for (index, field), issues in sorted(grouped.items()):
        item = items[index]
        original, current = _field_values(item, field)
        if not original:
            continue
        reasons = []
        for issue in issues:
            issue_id = issue.get("issue_id") or issue.get("type") or "translation_failure"
            reason = issue.get("reason") or ""
            label = f"{issue_id}: {reason}" if reason else issue_id
            if label not in reasons:
                reasons.append(label)
        entries.append({
            "item_id": stable_item_id(item, field, original),
            "source": item.get("source") or "",
            "date": str(item.get("date") or ""),
            "url": item.get("url") or "",
            "field": field,
            "original_text": original,
            "original_sha256": original_sha256(original),
            "title_original": item.get("title_original") or item.get("title") or "",
            "current_cn": current,
            "qa_failure": "\n".join(reasons),
        })
    return entries


def write_queue_workbook(entries, week_start, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = VISIBLE_SHEET
    sheet.merge_cells("A1:G1")
    sheet["A1"] = 'ONLY EDIT THE "Manual Translation" COLUMN. DO NOT MODIFY OTHER COLUMNS.'
    sheet["A1"].font = Font(bold=True, color="9C0006")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
    sheet.append(VISIBLE_HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.freeze_panes = "A3"
    widths = (24, 22, 12, 60, 50, 55, 60)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet["G2"].comment = Comment("Enter the complete approved Chinese translation here.", "DBG Weekly Intelligence")

    metadata = workbook.create_sheet(METADATA_SHEET)
    metadata.append(METADATA_HEADERS)
    for entry in entries:
        visible_row = sheet.max_row + 1
        sheet.append((
            entry["item_id"], entry["source"], entry["field"], entry["original_text"],
            entry["current_cn"], entry["qa_failure"], "",
        ))
        sheet.cell(visible_row, 7).fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in sheet[visible_row]:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment
        metadata.append((
            SCHEMA_VERSION, week_start, visible_row, entry["item_id"], entry["source"],
            entry["date"], entry["url"], entry["field"], entry["original_text"],
            entry["original_sha256"], entry["title_original"],
        ))
    metadata.sheet_state = "hidden"
    workbook.save(path)
    return path


def generate_queue_workbook(weekly_payload, qa_report, output_dir):
    entries = build_queue_entries(weekly_payload, qa_report)
    if not entries:
        return None
    compact = str(weekly_payload["week_start"]).replace("-", "")
    path = Path(output_dir) / f"manual_translation_queue_{compact}.xlsx"
    return write_queue_workbook(entries, str(weekly_payload["week_start"]), path)


def _read_workbook_rows(path, expected_week_start):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    if VISIBLE_SHEET not in workbook.sheetnames or METADATA_SHEET not in workbook.sheetnames:
        raise ManualTranslationError("Workbook is missing required Manual Review or _metadata sheet.")
    sheet = workbook[VISIBLE_SHEET]
    metadata = workbook[METADATA_SHEET]
    headers = tuple(cell.value for cell in metadata[1])
    if headers != METADATA_HEADERS:
        raise ManualTranslationError("Workbook metadata schema is invalid.")

    rows = []
    seen = set()
    missing = []
    for values in metadata.iter_rows(min_row=2, values_only=True):
        record = dict(zip(METADATA_HEADERS, values))
        if record["schema_version"] != SCHEMA_VERSION:
            raise ManualTranslationError(f"Unsupported workbook schema_version: {record['schema_version']}")
        if record["week_start"] != expected_week_start:
            raise ManualTranslationError(
                f"Workbook week_start mismatch: expected {expected_week_start}, got {record['week_start']}"
            )
        item_id = str(record["item_id"] or "")
        if not item_id or item_id in seen:
            raise ManualTranslationError(f"Duplicate or missing Item ID: {item_id or '<blank>'}")
        seen.add(item_id)
        if record["field"] not in {"title", "summary"}:
            raise ManualTranslationError(f"Unsupported translation field for Item ID {item_id}: {record['field']}")
        expected_hash = original_sha256(record["original_text"] or "")
        if record["original_sha256"] != expected_hash:
            raise ManualTranslationError(f"Original text hash mismatch for Item ID: {item_id}")
        identity = {
            "source": record["source"] or "",
            "date": str(record["date"] or ""),
            "url": record["url"] or "",
        }
        expected_item_id = stable_item_id(identity, record["field"], record["original_text"] or "")
        if item_id != expected_item_id:
            raise ManualTranslationError(f"Item ID metadata mismatch: {item_id}")
        visible_row = int(record["visible_row"])
        manual_translation = str(sheet.cell(visible_row, 7).value or "").strip()
        if not manual_translation:
            missing.append(item_id)
        record["manual_translation"] = manual_translation
        rows.append(record)
    if missing:
        raise ManualTranslationError("Missing Manual Translation for Item IDs: " + ", ".join(missing))
    return rows


def validate_manual_rows(rows):
    failures = []
    for row in rows:
        issues = [
            issue for issue in validate_translation_field(
                original=row["original_text"],
                translated=row["manual_translation"],
                field_type=row["field"],
                source=row["source"],
                title_original=row["title_original"],
            )
            if issue.get("severity") == FAIL
        ]
        if issues:
            failures.append({
                "item_id": row["item_id"],
                "source": row["source"],
                "field": row["field"],
                "issues": [{"issue_id": issue.get("issue_id"), "reason": issue.get("reason")} for issue in issues],
            })
    return failures


def approved_payload(rows, week_start):
    overrides = [{
        "item_id": row["item_id"],
        "source": row["source"],
        "date": str(row["date"] or ""),
        "url": row["url"] or "",
        "field": row["field"],
        "original_sha256": row["original_sha256"],
        "manual_translation": row["manual_translation"],
    } for row in sorted(rows, key=lambda value: value["item_id"])]
    return {"schema_version": SCHEMA_VERSION, "week_start": week_start, "overrides": overrides}


def atomic_write_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
            temp_path = Path(handle.name)
            json.dump(payload, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_path, path)
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink()


def apply_workbook(path, week_start, output_path):
    rows = _read_workbook_rows(path, week_start)
    failures = validate_manual_rows(rows)
    if failures:
        details = []
        for failure in failures:
            issue_text = ", ".join(
                f"{issue['issue_id']}: {issue['reason']}" for issue in failure["issues"]
            )
            details.append(
                f"{failure['item_id']} | {failure['source']} | {failure['field']} | {issue_text}"
            )
        raise ManualTranslationError("Manual translation QA failed:\n" + "\n".join(details))
    payload = approved_payload(rows, week_start)
    atomic_write_json(output_path, payload)
    return payload


def _clear_field_translation_issues(item, field):
    retained = []
    for issue in item.get("translation_quality_issues") or []:
        issue_field = issue.get("field") or (issue.get("evidence") or {}).get("field")
        if issue_field != field:
            retained.append(issue)
    if retained:
        item["translation_quality_issues"] = retained
    else:
        item.pop("translation_quality_issues", None)


def apply_approved_overrides_with_report(weekly_payload, approved_payload_data):
    report = {"applied": [], "absent": [], "identity_hash_conflict": [], "ambiguous": []}
    if not approved_payload_data:
        return report
    if approved_payload_data.get("schema_version") != SCHEMA_VERSION:
        report["identity_hash_conflict"] = list(approved_payload_data.get("overrides") or [])
        return report
    if approved_payload_data.get("week_start") != weekly_payload.get("week_start"):
        report["identity_hash_conflict"] = list(approved_payload_data.get("overrides") or [])
        return report

    override_rows = list(approved_payload_data.get("overrides") or [])
    overrides = {entry.get("item_id"): entry for entry in override_rows}
    matched = set()
    for item in weekly_payload.get("items") or []:
        for field in ("title", "summary"):
            original, _ = _field_values(item, field)
            if not original:
                continue
            item_id, override = _approved_override_for_item(overrides, item, field, original)
            if not override:
                continue
            identity_matches = (
                override.get("source") == (item.get("source") or "")
                and str(override.get("date") or "") == str(item.get("date") or "")
                and override.get("field") == field
                and override.get("original_sha256") == original_sha256(original)
                and (
                    (override.get("url") or "") == (item.get("url") or "")
                    or (
                        (item.get("source") or "") in INVESTMENT_ASSOCIATION_SOURCES
                        and _investment_association_canonical_url(override.get("url") or "")
                        == _investment_association_canonical_url(item.get("url") or "")
                        != ""
                    )
                )
            )
            if not identity_matches:
                continue
            item[f"{field}_cn"] = override.get("manual_translation") or ""
            provenance = dict(item.get("translation_provenance") or {})
            provenance[field] = "manual_approved"
            item["translation_provenance"] = provenance
            _clear_field_translation_issues(item, field)
            matched.add(item_id)
            report["applied"].append({
                "item_id": item_id,
                "source": override.get("source"),
                "date": override.get("date"),
                "field": field,
                "url": override.get("url"),
            })

    items = weekly_payload.get("items") or []
    for override in override_rows:
        if override.get("item_id") in matched:
            continue
        same_identity = []
        for item in items:
            if override.get("source") != (item.get("source") or ""):
                continue
            if str(override.get("date") or "") != str(item.get("date") or ""):
                continue
            field = override.get("field")
            original, _ = _field_values(item, field)
            if not original:
                continue
            same_identity.append((item, original))

        detail = {
            "item_id": override.get("item_id"),
            "source": override.get("source"),
            "date": override.get("date"),
            "field": override.get("field"),
            "url": override.get("url"),
        }
        if not same_identity:
            report["absent"].append(detail)
            continue
        hash_matches = [
            item for item, original in same_identity
            if original_sha256(original) == override.get("original_sha256")
        ]
        if not hash_matches:
            report["identity_hash_conflict"].append(detail)
            continue
        equivalent = [
            item for item in hash_matches
            if (item.get("url") or "") == (override.get("url") or "")
            or (
                (item.get("source") or "") in INVESTMENT_ASSOCIATION_SOURCES
                and _investment_association_canonical_url(item.get("url") or "")
                == _investment_association_canonical_url(override.get("url") or "")
                != ""
            )
        ]
        target_key = (
            override.get("source"), str(override.get("date") or ""), override.get("field"),
            override.get("original_sha256"), _investment_association_canonical_url(override.get("url") or ""),
        )
        equivalent_approvals = [
            row for row in override_rows
            if (
                row.get("source"), str(row.get("date") or ""), row.get("field"),
                row.get("original_sha256"), _investment_association_canonical_url(row.get("url") or ""),
            ) == target_key
        ]
        if len(equivalent) > 1 or len(equivalent_approvals) > 1:
            report["ambiguous"].append(detail)
        else:
            report["absent"].append(detail)
    return report


def apply_approved_overrides(weekly_payload, approved_payload_data):
    report = apply_approved_overrides_with_report(weekly_payload, approved_payload_data)
    applied = len(report["applied"])
    unmatched = sum(len(report[key]) for key in ("absent", "identity_hash_conflict", "ambiguous"))
    return applied, unmatched


def apply_approved_file(weekly_payload, week_start, approved_dir="manual_translations/approved"):
    compact = str(week_start).replace("-", "")
    path = Path(approved_dir) / f"weekly_{compact}.json"
    if not path.exists():
        return 0, 0
    with path.open(encoding="utf-8") as handle:
        approved = json.load(handle)
    return apply_approved_overrides(weekly_payload, approved)
