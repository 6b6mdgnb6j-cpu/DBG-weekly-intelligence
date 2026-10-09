import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from crawler.manual_translations import (
    METADATA_SHEET,
    VISIBLE_HEADERS,
    VISIBLE_SHEET,
    ManualTranslationError,
    apply_approved_file,
    apply_approved_overrides,
    apply_workbook,
    build_queue_entries,
    generate_queue_workbook,
    original_sha256,
    stable_item_id,
    write_queue_workbook,
)
from crawler.translation_quality import FAIL, audit_translation_quality, validate_translation_field


WEEK = "2026-09-14"
ISS_ORIGINAL = (
    "Institutional Shareholder Services (ISS) submitted a brief challenging an unconstitutional "
    "demand in a U.S. Securities & Exchange Commission subpoena."
)
ISS_MANUAL = "ISS提交简报，挑战SEC传票中的违宪要求。"


def item(**overrides):
    value = {
        "source": "ISS STOXX",
        "date": "2026-09-18",
        "url": "https://example.test/iss-brief",
        "title_original": "ISS challenges SEC demand",
        "title_cn": "ISS挑战SEC要求",
        "summary_original": ISS_ORIGINAL,
        "summary_cn": "ISS提出了违宪要求。",
    }
    value.update(overrides)
    return value


def issue(index=0, field="summary", issue_id="translation_relationship_missing", severity=FAIL):
    return {
        "index": index,
        "field": field,
        "issue_id": issue_id,
        "type": issue_id,
        "severity": severity,
        "reason": "Relationship direction is missing.",
    }


class ManualTranslationQueueTests(unittest.TestCase):
    def test_multiple_failures_for_same_field_create_one_aggregated_row(self):
        payload = {"week_start": WEEK, "items": [item()]}
        report = {"content_issues": [issue(), issue(issue_id="translation_entity_alignment_failure")]}

        entries = build_queue_entries(payload, report)

        self.assertEqual(len(entries), 1)
        self.assertIn("translation_relationship_missing", entries[0]["qa_failure"])
        self.assertIn("translation_entity_alignment_failure", entries[0]["qa_failure"])

    def test_title_and_summary_failures_create_two_rows(self):
        entries = build_queue_entries(
            {"week_start": WEEK, "items": [item()]},
            {"content_issues": [issue(field="title"), issue(field="summary")]},
        )
        self.assertEqual({entry["field"] for entry in entries}, {"title", "summary"})

    def test_warning_and_non_translation_fail_do_not_queue(self):
        entries = build_queue_entries(
            {"week_start": WEEK, "items": [item()]},
            {"content_issues": [
                issue(severity="WARNING"),
                issue(issue_id="outside_week_date"),
            ]},
        )
        self.assertEqual(entries, [])

    def test_stable_id_is_deterministic_and_changes_with_original(self):
        current = item()
        first = stable_item_id(current, "summary", ISS_ORIGINAL)
        self.assertEqual(first, stable_item_id(dict(current), "summary", ISS_ORIGINAL))
        self.assertNotEqual(first, stable_item_id(current, "summary", ISS_ORIGINAL + " Changed."))

    def test_workbook_has_visible_columns_and_hidden_authoritative_metadata(self):
        entries = build_queue_entries(
            {"week_start": WEEK, "items": [item()]}, {"content_issues": [issue()]}
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = write_queue_workbook(entries, WEEK, Path(tmp) / "queue.xlsx")
            workbook = load_workbook(path)
            sheet = workbook[VISIBLE_SHEET]
            metadata = workbook[METADATA_SHEET]

            self.assertEqual(tuple(cell.value for cell in sheet[2]), VISIBLE_HEADERS)
            self.assertEqual(sheet.cell(2, 7).value, "Manual Translation")
            self.assertEqual(metadata.sheet_state, "hidden")
            self.assertEqual(metadata.cell(2, 2).value, WEEK)

    def test_no_translation_failures_produce_no_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = generate_queue_workbook(
                {"week_start": WEEK, "items": [item()]}, {"content_issues": []}, tmp
            )
            self.assertIsNone(path)

    def test_repeated_character_candidate_reaches_final_qa_and_summary_queue(self):
        malformed = "ICE报告称，未平仓量上上上升了14个基点。"
        current = item(
            source="ICE",
            summary_original="ICE reports that open interest increased by 14 basis points during the period.",
            summary_cn=malformed,
            translation_quality_issues=[{
                "issue_id": "translation_sanity_quality_failure",
                "severity": FAIL,
                "evidence": {"reason": "repeated_character", "field": "summary", "source": "ICE"},
            }],
        )
        issues = audit_translation_quality([current], "bilingual")
        entries = build_queue_entries(
            {"week_start": WEEK, "items": [current]}, {"content_issues": issues}
        )

        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]["field"], "summary")
        self.assertEqual(entries[0]["current_cn"], malformed)
        self.assertIn("repeated_character", entries[0]["qa_failure"])

        with tempfile.TemporaryDirectory() as tmp:
            path = generate_queue_workbook(
                {"week_start": WEEK, "items": [current]},
                {"content_issues": issues},
                tmp,
            )
            workbook = load_workbook(path)
            visible = workbook[VISIBLE_SHEET]
            metadata = workbook[METADATA_SHEET]
            self.assertEqual(visible.cell(3, 1).value, entries[0]["item_id"])
            self.assertEqual(visible.cell(3, 2).value, "ICE")
            self.assertEqual(visible.cell(3, 3).value, "summary")
            self.assertEqual(visible.cell(3, 4).value, current["summary_original"])
            self.assertEqual(visible.cell(3, 5).value, malformed)
            self.assertIn("repeated_character", visible.cell(3, 6).value)
            self.assertIsNone(visible.cell(3, 7).value)
            self.assertEqual(metadata.sheet_state, "hidden")


class ManualTranslationCLITests(unittest.TestCase):
    def test_direct_script_invocation_imports_project_from_any_working_directory(self):
        script = Path(__file__).resolve().parents[1] / "scripts" / "apply_manual_translations.py"
        env = dict(os.environ)
        env.pop("PYTHONPATH", None)
        with tempfile.TemporaryDirectory() as tmp:
            for cwd in (script.parent.parent, Path(tmp)):
                with self.subTest(cwd=str(cwd)):
                    result = subprocess.run(
                        [
                            sys.executable,
                            str(script),
                            "--week-start",
                            "2026-09-14",
                            "--input",
                            str(Path(tmp) / "missing.xlsx"),
                        ],
                        cwd=cwd,
                        env=env,
                        capture_output=True,
                        text=True,
                        check=False,
                    )
                    self.assertEqual(result.returncode, 2)
                    self.assertIn("Manual translation workbook not found", result.stderr)
                    self.assertNotIn("ModuleNotFoundError", result.stderr)


class ManualTranslationApplyTests(unittest.TestCase):
    def make_workbook(self, root):
        entries = build_queue_entries(
            {"week_start": WEEK, "items": [item()]}, {"content_issues": [issue()]}
        )
        return write_queue_workbook(entries, WEEK, Path(root) / "queue.xlsx")

    def test_valid_workbook_writes_deterministic_approved_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            workbook = load_workbook(workbook_path)
            workbook[VISIBLE_SHEET].cell(3, 7).value = ISS_MANUAL
            workbook.save(workbook_path)
            output = Path(tmp) / "approved.json"

            payload = apply_workbook(workbook_path, WEEK, output)

            self.assertEqual(payload["schema_version"], 1)
            self.assertEqual(payload["overrides"][0]["manual_translation"], ISS_MANUAL)
            self.assertEqual(json.loads(output.read_text(encoding="utf-8")), payload)

    def test_blank_or_invalid_translation_does_not_mutate_approved_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            output = Path(tmp) / "approved.json"
            output.write_text('{"existing": true}\n', encoding="utf-8")
            with self.assertRaisesRegex(ManualTranslationError, "Missing Manual Translation"):
                apply_workbook(workbook_path, WEEK, output)
            self.assertEqual(output.read_text(encoding="utf-8"), '{"existing": true}\n')

            workbook = load_workbook(workbook_path)
            workbook[VISIBLE_SHEET].cell(3, 7).value = "ISS提出了违宪要求。"
            workbook.save(workbook_path)
            with self.assertRaisesRegex(ManualTranslationError, "translation_relationship_missing"):
                apply_workbook(workbook_path, WEEK, output)
            self.assertEqual(output.read_text(encoding="utf-8"), '{"existing": true}\n')

    def test_duplicate_ids_and_wrong_week_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            workbook = load_workbook(workbook_path)
            sheet = workbook[VISIBLE_SHEET]
            metadata = workbook[METADATA_SHEET]
            sheet.cell(3, 7).value = ISS_MANUAL
            metadata.append(tuple(cell.value for cell in metadata[2]))
            workbook.save(workbook_path)
            with self.assertRaisesRegex(ManualTranslationError, "Duplicate"):
                apply_workbook(workbook_path, WEEK, Path(tmp) / "approved.json")

            workbook = load_workbook(workbook_path)
            workbook.remove(workbook[METADATA_SHEET])
            metadata = workbook.create_sheet(METADATA_SHEET)
            from crawler.manual_translations import METADATA_HEADERS
            metadata.append(METADATA_HEADERS)
            entry = build_queue_entries(
                {"week_start": WEEK, "items": [item()]}, {"content_issues": [issue()]}
            )[0]
            metadata.append((1, "2026-09-21", 3, entry["item_id"], entry["source"], entry["date"],
                             entry["url"], entry["field"], entry["original_text"],
                             entry["original_sha256"], entry["title_original"]))
            metadata.sheet_state = "hidden"
            workbook.save(workbook_path)
            with self.assertRaisesRegex(ManualTranslationError, "week_start mismatch"):
                apply_workbook(workbook_path, WEEK, Path(tmp) / "approved.json")

    def test_visible_metadata_tampering_is_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            workbook = load_workbook(workbook_path)
            sheet = workbook[VISIBLE_SHEET]
            sheet.cell(3, 1).value = "tampered-id"
            sheet.cell(3, 2).value = "Tampered Source"
            sheet.cell(3, 4).value = "Tampered original"
            sheet.cell(3, 7).value = ISS_MANUAL
            workbook.save(workbook_path)

            payload = apply_workbook(workbook_path, WEEK, Path(tmp) / "approved.json")

            self.assertEqual(payload["overrides"][0]["source"], "ISS STOXX")
            self.assertEqual(payload["overrides"][0]["original_sha256"], original_sha256(ISS_ORIGINAL))

    def test_hidden_original_hash_tampering_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            workbook = load_workbook(workbook_path)
            workbook[VISIBLE_SHEET].cell(3, 7).value = ISS_MANUAL
            workbook[METADATA_SHEET].cell(2, 10).value = "0" * 64
            workbook.save(workbook_path)

            with self.assertRaisesRegex(ManualTranslationError, "Original text hash mismatch"):
                apply_workbook(workbook_path, WEEK, Path(tmp) / "approved.json")

    def test_approved_json_uses_atomic_replace(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = self.make_workbook(tmp)
            workbook = load_workbook(workbook_path)
            workbook[VISIBLE_SHEET].cell(3, 7).value = ISS_MANUAL
            workbook.save(workbook_path)
            output = Path(tmp) / "approved.json"

            from crawler import manual_translations
            real_replace = manual_translations.os.replace
            calls = []

            def recording_replace(source, target):
                calls.append((Path(source), Path(target)))
                return real_replace(source, target)

            with patch.object(manual_translations.os, "replace", side_effect=recording_replace):
                apply_workbook(workbook_path, WEEK, output)

            self.assertEqual(len(calls), 1)
            self.assertEqual(calls[0][1], output)
            self.assertEqual(calls[0][0].parent, output.parent)
            self.assertTrue(output.exists())


class ManualTranslationOverrideTests(unittest.TestCase):
    def approved(self, current_item, field="summary", translation=ISS_MANUAL, **changes):
        original = current_item[f"{field}_original"]
        override = {
            "item_id": stable_item_id(current_item, field, original),
            "source": current_item["source"],
            "date": current_item["date"],
            "url": current_item["url"],
            "field": field,
            "original_sha256": original_sha256(original),
            "manual_translation": translation,
        }
        override.update(changes)
        return {"schema_version": 1, "week_start": WEEK, "overrides": [override]}

    def test_exact_override_applies_records_provenance_and_clears_only_same_field_issue(self):
        current = item(translation_quality_issues=[
            {"issue_id": "translation_relationship_missing", "evidence": {"field": "summary"}},
            {"issue_id": "translation_title_proposition_missing", "evidence": {"field": "title"}},
        ])
        payload = {"week_start": WEEK, "items": [current]}

        applied, stale = apply_approved_overrides(payload, self.approved(current))

        self.assertEqual((applied, stale), (1, 0))
        self.assertEqual(current["summary_cn"], ISS_MANUAL)
        self.assertEqual(current["translation_provenance"]["summary"], "manual_approved")
        self.assertEqual(len(current["translation_quality_issues"]), 1)
        self.assertEqual(current["translation_quality_issues"][0]["evidence"]["field"], "title")

    def test_identity_mismatches_do_not_apply_or_clear_fresh_failure(self):
        for changed in (
            {"original_sha256": "0" * 64},
            {"url": "https://example.test/other"},
            {"date": "2026-09-19"},
            {"source": "Other"},
        ):
            with self.subTest(changed=changed):
                current = item(translation_quality_issues=[
                    {"issue_id": "translation_relationship_missing", "evidence": {"field": "summary"}}
                ])
                payload = {"week_start": WEEK, "items": [current]}
                applied, stale = apply_approved_overrides(payload, self.approved(current, **changed))
                self.assertEqual((applied, stale), (0, 1))
                self.assertEqual(current["summary_cn"], "ISS提出了违宪要求。")
                self.assertTrue(current["translation_quality_issues"])

    def test_investment_association_legacy_url_variants_match_title_and_summary(self):
        old_url = (
            "https://www.theia.org/news/press-releases/"
            "summer-inflows-continue-investors-place-ps894-million-funds-august"
        )
        new_url = (
            "https://www.theia.org/index%2Ephp/news/press-releases/"
            "summer-inflows-continue-investors-place-ps894-million-funds-august"
        )
        current = item(
            source="英国投资管理协会",
            date="2026-10-01",
            url=new_url,
            title_original="Summer inflows continue as investors place £894 million into funds in August",
            title_cn="错误标题",
            summary_original=(
                "1 October 2026, London – Net retail sales were £894 million in August, marking the "
                "tenth consecutive month of retail inflows according to recent data from the Investment "
                "Association (IA). This marks a notable change from August 2025, where outflows reached "
                "£1.8 billion, underscoring the improvement"
            ),
            summary_cn="错误摘要",
            translation_quality_issues=[
                {"issue_id": "translation_monetary_mismatch", "evidence": {"field": "title"}},
                {"issue_id": "translation_monetary_mismatch", "evidence": {"field": "summary"}},
            ],
        )
        approved_item = dict(current, url=old_url)
        overrides = []
        for field, translation in (("title", "8月投资者向基金投入8.94亿英镑"), ("summary", "8月投资者向基金投入8.94亿英镑。")):
            original = approved_item[f"{field}_original"]
            overrides.append({
                "item_id": stable_item_id(approved_item, field, original),
                "source": approved_item["source"],
                "date": approved_item["date"],
                "url": old_url,
                "field": field,
                "original_sha256": original_sha256(original),
                "manual_translation": translation,
            })

        applied, stale = apply_approved_overrides(
            {"week_start": WEEK, "items": [current]},
            {"schema_version": 1, "week_start": WEEK, "overrides": overrides},
        )

        self.assertEqual((applied, stale), (2, 0))
        self.assertEqual(current["title_cn"], "8月投资者向基金投入8.94亿英镑")
        self.assertEqual(current["summary_cn"], "8月投资者向基金投入8.94亿英镑。")
        self.assertEqual(current["translation_provenance"], {
            "title": "manual_approved",
            "summary": "manual_approved",
        })
        self.assertNotIn("translation_quality_issues", current)

    def test_investment_association_equivalence_is_bidirectional(self):
        encoded = "https://www.theia.org/index%2ephp/news/press-releases/article"
        canonical = "https://www.theia.org/news/press-releases/article"
        current = item(source="Investment Association", url=canonical)
        approved = self.approved(current, url=encoded)
        approved["overrides"][0]["item_id"] = stable_item_id(
            dict(current, url=encoded), "summary", current["summary_original"]
        )

        self.assertEqual(
            apply_approved_overrides({"week_start": WEEK, "items": [current]}, approved),
            (1, 0),
        )

    def test_investment_association_unrelated_url_and_hash_mismatch_do_not_match(self):
        old_url = "https://www.theia.org/news/press-releases/article-one"
        variants = (
            {"url": "https://www.theia.org/index%2Ephp/news/press-releases/article-two"},
            {"url": "https://www.theia.org/index%2Ephp/news/press-releases/article-one", "original_sha256": "0" * 64},
            {"url": "https://other.example/news/press-releases/article-one"},
        )
        for changes in variants:
            with self.subTest(changes=changes):
                current = item(source="英国投资管理协会", url=changes["url"])
                approved_item = item(source="英国投资管理协会", url=old_url)
                approved = self.approved(approved_item, **{
                    key: value for key, value in changes.items() if key != "url"
                })

                self.assertEqual(
                    apply_approved_overrides({"week_start": WEEK, "items": [current]}, approved),
                    (0, 1),
                )

    def test_investment_association_ambiguous_equivalent_overrides_are_rejected(self):
        canonical = "https://www.theia.org/news/press-releases/article"
        encoded = "https://www.theia.org/index%2Ephp/news/press-releases/article"
        current = item(source="英国投资管理协会", url=encoded)
        first = self.approved(dict(current, url=canonical))["overrides"][0]
        second = dict(first, item_id="legacy-conflicting-id", manual_translation="另一人工译文")

        applied, stale = apply_approved_overrides(
            {"week_start": WEEK, "items": [current]},
            {"schema_version": 1, "week_start": WEEK, "overrides": [first, second]},
        )

        self.assertEqual((applied, stale), (0, 2))

    def test_no_approved_file_leaves_payload_unchanged(self):
        with tempfile.TemporaryDirectory() as tmp:
            payload = {"week_start": WEEK, "items": [item()]}
            before = json.dumps(payload, ensure_ascii=False, sort_keys=True)
            self.assertEqual(apply_approved_file(payload, WEEK, tmp), (0, 0))
            self.assertEqual(json.dumps(payload, ensure_ascii=False, sort_keys=True), before)

    def test_final_qa_accepts_good_manual_and_rejects_bad_manual(self):
        good = validate_translation_field(ISS_ORIGINAL, ISS_MANUAL, "summary", "ISS STOXX")
        bad = validate_translation_field(ISS_ORIGINAL, "ISS提出了违宪要求。", "summary", "ISS STOXX")
        self.assertFalse(any(value.get("severity") == FAIL for value in good))
        self.assertIn("translation_relationship_missing", [value.get("issue_id") for value in bad])

    def test_manual_validation_rejects_fatal_sanity_corruption(self):
        cases = {
            "repeated_character": "市场上上上升了14个基点。",
            "repeated_phrase": "市场下跌市场下跌市场下跌市场下跌市场下跌。",
            "replacement_character": "市场\ufffd变化。",
            "control_ratio": "市场\x00\x01变化。",
        }
        for reason, translated in cases.items():
            with self.subTest(reason=reason):
                issues = validate_translation_field(
                    "Market conditions changed during the reporting period.",
                    translated,
                    "summary",
                    "Example",
                )
                sanity = [
                    value for value in issues
                    if value.get("issue_id") == "translation_sanity_quality_failure"
                ]
                self.assertTrue(sanity)
                self.assertEqual(sanity[0]["evidence"]["sanity_reason"], reason)

    def test_exact_iss_failure_is_resolved_without_translation_code_change(self):
        current = item(translation_quality_issues=[
            {"issue_id": "translation_relationship_missing", "evidence": {"field": "summary"}}
        ])
        payload = {"week_start": WEEK, "items": [current]}

        applied, stale = apply_approved_overrides(payload, self.approved(current))
        final_issues = validate_translation_field(
            ISS_ORIGINAL, current["summary_cn"], "summary", current["source"]
        )

        self.assertEqual((applied, stale), (1, 0))
        self.assertFalse(any(value.get("severity") == FAIL for value in final_issues))

    def test_three_approved_overrides_apply_while_new_sanity_failure_stays_queued(self):
        approved_items = []
        overrides = []
        for index in range(3):
            current = item(
                url=f"https://example.test/approved-{index}",
                translation_quality_issues=[{
                    "issue_id": "translation_relationship_missing",
                    "evidence": {"field": "summary"},
                }],
            )
            approved_items.append(current)
            overrides.append(self.approved(current)["overrides"][0])

        unresolved = item(
            source="ICE",
            url="https://example.test/new-ice-failure",
            summary_original="ICE reports that open interest increased by 14 basis points during the period.",
            summary_cn="ICE报告称，未平仓量上上上升了14个基点。",
            translation_quality_issues=[{
                "issue_id": "translation_sanity_quality_failure",
                "severity": FAIL,
                "evidence": {"reason": "repeated_character", "field": "summary", "source": "ICE"},
            }],
        )
        payload = {"week_start": WEEK, "items": approved_items + [unresolved]}

        applied, stale = apply_approved_overrides(
            payload,
            {"schema_version": 1, "week_start": WEEK, "overrides": overrides},
        )
        issues = audit_translation_quality(payload["items"], "bilingual")
        entries = build_queue_entries(payload, {"content_issues": issues})

        self.assertEqual((applied, stale), (3, 0))
        self.assertTrue(all(
            value.get("translation_provenance", {}).get("summary") == "manual_approved"
            for value in approved_items
        ))
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]["source"], "ICE")
        self.assertEqual(entries[0]["field"], "summary")
        self.assertIn("repeated_character", entries[0]["qa_failure"])


if __name__ == "__main__":
    unittest.main()
